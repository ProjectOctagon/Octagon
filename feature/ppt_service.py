from pptx import Presentation
from pptx.util import Pt, Cm
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.text import MSO_ANCHOR
import os
from datetime import datetime
from feature.slide_models import SlideModel, BoxModel, PresentationModel

class PptService:

    COLORS = {
        "HEAD": RGBColor(0, 15, 45),       # dark blue
        "ASSISTANT": RGBColor(26, 153, 171),# teal
        "UNIT": RGBColor(189, 215, 255),    # light Blue
        "DOTTED": RGBColor(89, 89, 89),     # grey
        "HIGHLIGHT": RGBColor(0, 204, 102)  # green for change of roles
    }


    def __init__(self, template_dir: str, output_dir: str, base_url: str):
        self.template_dir = template_dir # where .pptx templates for generating the slides are stored
        self.output_dir = output_dir # path for saving generated files
        self.base_url = base_url #  used to generate a download link

    def generate_presentation(self, model, template_name: str, output_name: str = None): # model => list of slide definitions, template_name =>.pptx template used to generate slides, output_name => optional filename for generated pptx                                                                                
        template_path = os.path.join(self.template_dir, f"{template_name}.pptx") #loading template for generating the slides

        if not os.path.exists(template_path): #if the template doesn't exist, throws an error
            raise FileNotFoundError(f"Template not found: {template_path}")

        prs = Presentation(template_path)

        #remove the placeholder slide from the template
        for i in range(len(prs.slides) - 1, -1, -1):
            rId = prs.slides._sldIdLst[i].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[i]

        # build new slides with the requred structure
        for slide_model in model.slides: #model contains a list of slide definitions;
            slide_type = getattr(slide_model, "type") #get type required for slides and generate appropriate pptx

            if slide_type == "org_chart":  
                self._add_org_chart_slide(prs, slide_model)
            

        # save file
        filename = self._generate_filename(output_name) 
        output_path = os.path.join(self.output_dir, filename)

        os.makedirs(self.output_dir, exist_ok=True)
        prs.save(output_path)

        return self._build_file_url(filename)

    
    def _add_org_chart_slide(self, prs, slide_model: SlideModel):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        #remove any existing "Click to add text" boxes
        for shape in list(slide.shapes): 
            if shape.is_placeholder:
                sp = shape._sp
                sp.getparent().remove(sp)

   
        title_shape = slide.shapes.add_textbox(Cm(1.3), Cm(0.76), Cm(22.86), Cm(2.54))
        tf = title_shape.text_frame
        tf.text = f"{slide_model.title} {slide_model.hierarchy_level}"
        tf.paragraphs[0].font.size = Pt(28)
        tf.paragraphs[0].font.color.rgb = RGBColor(0, 102, 255)

       
        col_width = Cm(9)
        start_y = Cm(4.5)
        
        for i, column in enumerate(slide_model.columns):
            current_x = Cm(3.8) + (i * (col_width + Cm(0.76)))
            for j, unit in enumerate(column):
                unit_y = start_y + (j * Cm(2.16))
                self._draw_box(slide, current_x, unit_y, unit)

      
        self._draw_box(slide, Cm(7.62), Cm(15.75), slide_model.head_of_unit, width=Cm(6.93))
        if slide_model.assistant:
            self._draw_box(slide, Cm(14.23), Cm(15.75), slide_model.assistant, width=Cm(6.93))

        self._add_footer(slide)


    def _add_footer(self, slide): #adds a standard footer to the bottom of the slide
        left = Cm(5.57)
        top = Cm(17.24) 
        width = Cm(22.80)
        height = Cm(1.54) 
        
        footer_shape = slide.shapes.add_textbox(left, top, width, height)
        tf = footer_shape.text_frame
        tf.word_wrap = True 
        
        p = tf.paragraphs[0]
        p.text = "This document aims to support the business groups and corporate functions in Nokia subject to necessary legal procedures and approvals.  Nothing herein should be deemed to indicate that any final decision has been made which otherwise may require information and/or consultation with relevant employee representative body/ies, where applicable.​" 
        p.font.size = Pt(10.7) 
        p.font.color.rgb = RGBColor(128, 128, 128)
        p.alignment = PP_ALIGN.LEFT
    
    def _draw_box(self, slide, x, y, box_data: BoxModel, width=Cm(6.93), height=Cm(1.33)):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, height)
        
      
        tf = shape.text_frame 
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

       
        if box_data.is_changed:
            shape.line.color.rgb = self.COLORS["HIGHLIGHT"]
            shape.line.width = Pt(2.5)
        else:
            shape.line.fill.background()

       
        shape.fill.solid()
        if box_data.box_type == "manager":
            shape.fill.fore_color.rgb = self.COLORS["HEAD"] 
        elif box_data.box_type == "assistant":
            shape.fill.fore_color.rgb = self.COLORS["ASSISTANT"]
        elif box_data.box_type == "dotted":
            shape.fill.fore_color.rgb = self.COLORS["DOTTED"]
        else:
            shape.fill.fore_color.rgb = self.COLORS["UNIT"]

    
        tf.margin_left = Cm(0.25)
        
        p1 = tf.paragraphs[0]
        p1.text = box_data.title
        p1.font.bold = True
        p1.font.size = Pt(10)
        
     
        if box_data.box_type in ["manager", "assistant", "dotted"]:
            p1.font.color.rgb = RGBColor(255, 255, 255)
        else:
            p1.font.color.rgb = RGBColor(0, 0, 0)

        p2 = tf.add_paragraph()
        p2.text = box_data.subtitle
        p2.font.size = Pt(9)
        p2.font.color.rgb = p1.font.color.rgb


    
    def _generate_filename(self, output_name): #generates unique file names using timestamp to prevent overwriting existing files
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")  

        if output_name:
            return f"{output_name}_{timestamp}.pptx" 

        return f"presentation_{timestamp}.pptx" #if user doesn't provide name, it gives it a default one, example: presentation_20260410_153012.pptx

    def _build_file_url(self, filename): 
        return f"{self.base_url}/{filename}"